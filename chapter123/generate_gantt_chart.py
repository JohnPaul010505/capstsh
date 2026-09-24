# -*- coding: utf-8 -*-
"""
Generate the Gantt Chart (Work Plan) figure for the Triple J Fitness Center
capstone document - Chapter 3.11.

Layout follows the reference "Gantt chart" sample: an Activities column on
the left and one column per month (Jan..Dec); months covered by an activity
are shaded blue. There is NO quarter grouping row.

Everything is driven by the ACTIVITIES list below - edit a month span and
re-run to regenerate every output:

  Gantt-Fig29-Work-Plan.drawio      draw.io figure (repo .drawio convention)
  Gantt-Chart-Fig29.docx            Word table (Table Grid), paste-ready
  Gantt-Chart-Preview.png           Pillow render for quick preview
  Gantt-Captions-and-Narratives.md  caption + narrative markdown

Usage:
  python generate_gantt_chart.py           generate all outputs
  python generate_gantt_chart.py verify    generate, then verify all outputs
"""

import os
import sys
import xml.etree.ElementTree as ET
from xml.sax.saxutils import escape

# ---------------------------------------------------------------------------
# Config - edit here only
# ---------------------------------------------------------------------------

SECTION_TITLE = "3.11 GANTT CHART"
FIG_CAPTION = "Figure 29: Gantt Chart (Work Plan)"
CHART_TITLE = "Gantt chart"

MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
          "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]

# (activity label, [months covered])
ACTIVITIES = [
    ("Planning", ["Jan", "Feb"]),
    ("Data Gathering (Observation & Face-to-Face Interview)", ["Feb", "Mar"]),
    ("Group Meeting", ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug"]),
    ("Identification of appropriate software to use", ["Apr"]),
    ("Conceptualise Database", ["May"]),
    ("Database Design", ["May"]),
    ("Created Design for UI", ["Jun", "Jul"]),
    ("Developing System", ["Jun", "Jul", "Aug"]),
    ("System Testing", ["Aug"]),
    ("Implementing Changes", ["Aug"]),
    ("Beta Testing", ["Aug"]),
    ("Implementing Feedback", ["Sep"]),
    ("Consultation", ["Sep"]),
    ("Finishing Documentation", ["Sep"]),
]

FONT = "Times New Roman"
BAR_FILL = "8EAADB"    # blue bars (Word "Blue, Accent 1, Lighter 40%")
BAND_FILL = "B4C7E7"   # light blue title band

OUT_DIR = os.path.dirname(os.path.abspath(__file__))
DRAWIO_PATH = os.path.join(OUT_DIR, "Gantt-Fig29-Work-Plan.drawio")
DOCX_PATH = os.path.join(OUT_DIR, "Gantt-Chart-Fig29.docx")
PNG_PATH = os.path.join(OUT_DIR, "Gantt-Chart-Preview.png")
MD_PATH = os.path.join(OUT_DIR, "Gantt-Captions-and-Narratives.md")

MONTH_INDEX = {m: i for i, m in enumerate(MONTHS)}


def active_cell_count():
    return sum(len(ms) for _, ms in ACTIVITIES)


# ---------------------------------------------------------------------------
# draw.io output
# ---------------------------------------------------------------------------

LBL_W, MON_W = 200, 44
HEAD_H, ROW_H, TITLE_H = 28, 34, 32
MARGIN = 40
TABLE_W = LBL_W + len(MONTHS) * MON_W
TITLE_W = 320

S_TITLE = ("rounded=0;whiteSpace=wrap;html=1;fillColor=#%s;strokeColor=#000000;"
           "fontFamily=%s;fontSize=14;fontStyle=1;" % (BAND_FILL, FONT))
S_HEAD = ("rounded=0;whiteSpace=wrap;html=1;fillColor=#FFFFFF;strokeColor=#000000;"
          "fontFamily=%s;fontSize=12;fontStyle=1;" % FONT)
S_LABEL = ("rounded=0;whiteSpace=wrap;html=1;fillColor=#FFFFFF;strokeColor=#000000;"
           "align=left;spacingLeft=4;fontFamily=%s;fontSize=11;" % FONT)
S_EMPTY = "rounded=0;whiteSpace=wrap;html=1;fillColor=#FFFFFF;strokeColor=#000000;"
S_BAR = "rounded=0;whiteSpace=wrap;html=1;fillColor=#%s;strokeColor=#000000;" % BAR_FILL


def _vtx(cid, style, value, x, y, w, h):
    return ('<mxCell id="%s" value="%s" style="%s" vertex="1" parent="1">'
            '<mxGeometry x="%d" y="%d" width="%d" height="%d" as="geometry" />'
            '</mxCell>' % (cid, escape(value), style, x, y, w, h))


def build_drawio(path=DRAWIO_PATH):
    y0 = MARGIN + TITLE_H + 12
    cells = []

    cells.append(_vtx("title", S_TITLE, CHART_TITLE,
                      MARGIN + (TABLE_W - TITLE_W) // 2, MARGIN, TITLE_W, TITLE_H))
    cells.append(_vtx("h_lbl", S_HEAD, "Activities", MARGIN, y0, LBL_W, HEAD_H))
    for j, m in enumerate(MONTHS):
        cells.append(_vtx("h_%d" % j, S_HEAD, m,
                          MARGIN + LBL_W + j * MON_W, y0, MON_W, HEAD_H))

    for i, (label, months) in enumerate(ACTIVITIES):
        y = y0 + HEAD_H + i * ROW_H
        cells.append(_vtx("r%d_lbl" % i, S_LABEL, label, MARGIN, y, LBL_W, ROW_H))
        for j, m in enumerate(MONTHS):
            style = S_BAR if m in months else S_EMPTY
            cells.append(_vtx("r%d_c%d" % (i, j), style, "",
                              MARGIN + LBL_W + j * MON_W, y, MON_W, ROW_H))

    total_h = y0 + HEAD_H + len(ACTIVITIES) * ROW_H + MARGIN
    xml = [
        '<mxfile host="app.diagrams.net">',
        '  <diagram id="gantt29" name="Gantt Chart">',
        '    <mxGraphModel dx="1000" dy="700" grid="0" gridSize="10" guides="1" '
        'tooltips="1" connect="1" arrows="1" fold="1" page="1" pageScale="1" '
        'pageWidth="850" pageHeight="%d" math="0" shadow="0">' % max(1100, total_h),
        '      <root>',
        '        <mxCell id="0" />',
        '        <mxCell id="1" parent="0" />',
    ]
    xml.extend("        " + c for c in cells)
    xml += [
        '      </root>',
        '    </mxGraphModel>',
        '  </diagram>',
        '</mxfile>',
        '',
    ]
    with open(path, 'w', encoding='utf-8') as f:
        f.write("\n".join(xml))
    return path

# ---------------------------------------------------------------------------
# Word output (python-docx)
# ---------------------------------------------------------------------------

def build_docx(path=DOCX_PATH):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import (WD_TABLE_ALIGNMENT,
                                 WD_CELL_VERTICAL_ALIGNMENT,
                                 WD_ROW_HEIGHT_RULE)
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    def shade(cell_obj, fill):
        tc_pr = cell_obj._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), fill)
        tc_pr.append(shd)

    def put(cell_obj, text, bold=False, center=True, size=10):
        cell_obj.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell_obj.paragraphs[0]
        p.alignment = (WD_ALIGN_PARAGRAPH.CENTER if center
                       else WD_ALIGN_PARAGRAPH.LEFT)
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(1)
        run = p.add_run(text)
        run.font.name = FONT
        run.font.size = Pt(size)
        run.bold = bold

    doc = Document()
    doc.styles["Normal"].font.name = FONT
    doc.styles["Normal"].font.size = Pt(11)

    head = doc.add_heading(SECTION_TITLE, level=2)
    for run in head.runs:
        run.font.name = FONT
        run.font.color.rgb = RGBColor(0, 0, 0)

    # light blue "Gantt chart" title band, centered above the table
    band = doc.add_table(rows=1, cols=1)
    band.style = doc.styles["Table Grid"]
    band.alignment = WD_TABLE_ALIGNMENT.CENTER
    band.columns[0].width = Inches(3.0)
    bc = band.cell(0, 0)
    bc.width = Inches(3.0)
    shade(bc, BAND_FILL)
    put(bc, CHART_TITLE, bold=True, center=True, size=12)
    doc.add_paragraph()

    tbl = doc.add_table(rows=1 + len(ACTIVITIES), cols=1 + len(MONTHS))
    tbl.style = doc.styles["Table Grid"]
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False

    lbl_w, mon_w = Inches(1.9), Inches(0.38)

    hdr = tbl.rows[0]
    hdr.height = Inches(0.3)
    hdr.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    c = tbl.cell(0, 0)
    c.width = lbl_w
    put(c, "Activities", bold=True)
    for j, m in enumerate(MONTHS):
        c = tbl.cell(0, j + 1)
        c.width = mon_w
        put(c, m, bold=True)

    for i, (label, months) in enumerate(ACTIVITIES, start=1):
        row = tbl.rows[i]
        row.height = Inches(0.26)
        row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        c = tbl.cell(i, 0)
        c.width = lbl_w
        put(c, label, center=False, size=9)
        for j, m in enumerate(MONTHS):
            c = tbl.cell(i, j + 1)
            c.width = mon_w
            put(c, "")
            if m in months:
                shade(c, BAR_FILL)

    doc.save(path)
    return path
# ---------------------------------------------------------------------------
# PNG preview (Pillow)
# ---------------------------------------------------------------------------

PNG_LBL_W, PNG_MON_W = 230, 46
PNG_HEAD_H, PNG_ROW_H, PNG_TITLE_H = 30, 36, 34
PNG_MARGIN = 20
PNG_BAND_W = 300


def _font(bold=False, size=12):
    from PIL import ImageFont
    names = ("timesbd.ttf", "arialbd.ttf") if bold else ("times.ttf", "arial.ttf")
    for n in names:
        try:
            return ImageFont.truetype(n, size)
        except OSError:
            continue
    try:
        return ImageFont.load_default(size=size)
    except TypeError:
        return ImageFont.load_default()


def _wrap(d, text, font, max_w):
    words, lines, cur = text.split(), [], ""
    for w in words:
        t = (cur + " " + w).strip()
        if not cur or d.textlength(t, font=font) <= max_w:
            cur = t
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def png_cell_center(row, month):
    """Centre pixel of data cell (row, month) in the PNG - used by verify_png."""
    x0 = PNG_MARGIN + PNG_LBL_W + MONTH_INDEX[month] * PNG_MON_W + PNG_MON_W / 2.0
    y0 = (PNG_MARGIN + PNG_TITLE_H + 10 + PNG_HEAD_H
          + row * PNG_ROW_H + PNG_ROW_H / 2.0)
    return x0, y0


def build_png(path=PNG_PATH):
    from PIL import Image, ImageDraw

    W = PNG_MARGIN * 2 + PNG_LBL_W + len(MONTHS) * PNG_MON_W
    H = PNG_MARGIN * 2 + PNG_TITLE_H + 10 + PNG_HEAD_H + len(ACTIVITIES) * PNG_ROW_H
    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)

    def rect(x, y, w, h, fill):
        d.rectangle([x, y, x + w, y + h], fill=fill, outline="black")

    # title band
    bx = (W - PNG_BAND_W) // 2
    by = PNG_MARGIN
    rect(bx, by, PNG_BAND_W, PNG_TITLE_H, "#" + BAND_FILL)
    f_title = _font(bold=True, size=15)
    tw = d.textlength(CHART_TITLE, font=f_title)
    d.text((W / 2.0 - tw / 2, by + (PNG_TITLE_H - 16) / 2.0),
           CHART_TITLE, font=f_title, fill="black")

    x0, y0 = PNG_MARGIN, PNG_MARGIN + PNG_TITLE_H + 10

    # header row
    f_head = _font(bold=True, size=12)
    rect(x0, y0, PNG_LBL_W, PNG_HEAD_H, "white")
    tw = d.textlength("Activities", font=f_head)
    d.text((x0 + (PNG_LBL_W - tw) / 2.0, y0 + (PNG_HEAD_H - 13) / 2.0),
           "Activities", font=f_head, fill="black")
    for j, m in enumerate(MONTHS):
        x = x0 + PNG_LBL_W + j * PNG_MON_W
        rect(x, y0, PNG_MON_W, PNG_HEAD_H, "white")
        tw = d.textlength(m, font=f_head)
        d.text((x + (PNG_MON_W - tw) / 2.0, y0 + (PNG_HEAD_H - 13) / 2.0),
               m, font=f_head, fill="black")

    # data rows
    f_lbl = _font(size=11)
    f_small = _font(size=9)
    bar = "#" + BAR_FILL
    for i, (label, months) in enumerate(ACTIVITIES):
        y = y0 + PNG_HEAD_H + i * PNG_ROW_H
        rect(x0, y, PNG_LBL_W, PNG_ROW_H, "white")
        lines = _wrap(d, label, f_lbl, PNG_LBL_W - 8)
        f_use, lh = f_lbl, 13
        if len(lines) * lh > PNG_ROW_H - 6:
            f_use, lh = f_small, 11
            lines = _wrap(d, label, f_use, PNG_LBL_W - 8)
        ty = y + (PNG_ROW_H - len(lines) * lh) / 2.0
        for ln in lines:
            d.text((x0 + 4, ty), ln, font=f_use, fill="black")
            ty += lh
        for j, m in enumerate(MONTHS):
            x = x0 + PNG_LBL_W + j * PNG_MON_W
            rect(x, y, PNG_MON_W, PNG_ROW_H, bar if m in months else "white")

    img.save(path)
    return path

# ---------------------------------------------------------------------------
# Markdown captions and narratives
# ---------------------------------------------------------------------------

MONTH_FULL = {"Jan": "January", "Feb": "February", "Mar": "March",
              "Apr": "April", "May": "May", "Jun": "June", "Jul": "July",
              "Aug": "August", "Sep": "September", "Oct": "October",
              "Nov": "November", "Dec": "December"}


def _span(months):
    if len(months) == 1:
        return MONTH_FULL[months[0]]
    return "%s to %s" % (MONTH_FULL[months[0]], MONTH_FULL[months[-1]])


NARRATIVE = """Figure 29 presents the Gantt chart of the development of the proposed
Intelligent Fitness Progress Monitoring Application of Triple J Fitness Center.
The chart covers the months of January to September and follows the four phases
of the Rapid Application Development (RAD) methodology discussed in Section 3.1.
The proponents began with Planning in {planning}, followed by Data Gathering
through observation and face-to-face interviews in {data_gathering}. Group
meetings were conducted regularly from {group_meeting} to keep the development
of the system coordinated. The appropriate software to use was identified in
{software}, the database was conceptualized and designed in {database}, and the
user interface was designed in {ui}. The development of the system ran
through {development}, after which system testing, the implementation of
changes, and beta testing were carried out in {testing}. The feedback gathered
from the beta testers was implemented in {feedback}, and consultation with the
adviser together with the finishing of the documentation were completed in
{documentation}, marking the end of the planned development activities.""".format(
    planning=_span(ACTIVITIES[0][1]),
    data_gathering=_span(ACTIVITIES[1][1]),
    group_meeting=_span(ACTIVITIES[2][1]),
    software=_span(ACTIVITIES[3][1]),
    database=_span(ACTIVITIES[4][1]),
    ui=_span(ACTIVITIES[6][1]),
    development=_span(ACTIVITIES[7][1]),
    testing=_span(ACTIVITIES[8][1]),
    feedback=_span(ACTIVITIES[11][1]),
    documentation=_span(ACTIVITIES[13][1]),
)


def build_md(path=MD_PATH):
    lines = [
        "# 3.11 GANTT CHART — Captions and Narratives",
        "",
        "Caption and narrative paragraph for the Gantt chart (work plan) figure,",
        "generated by `generate_gantt_chart.py`. Figure numbering continues from the",
        "Use Case Diagram figures (Figures 26–28).",
        "",
        "---",
        "",
        "## Figure 29",
        "",
        "**Caption:** " + FIG_CAPTION,
        "",
        "**Narrative:**",
        "",
    ]
    lines.extend("> " + ln if ln else ">" for ln in NARRATIVE.splitlines())
    lines.append("")
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return path

# ---------------------------------------------------------------------------
# Verification
# ---------------------------------------------------------------------------

def _cell_fill(cell_obj):
    from docx.oxml.ns import qn
    tc_pr = cell_obj._tc.tcPr
    if tc_pr is None:
        return None
    shd = tc_pr.find(qn("w:shd"))
    return shd.get(qn("w:fill")) if shd is not None else None


def verify_drawio(path=DRAWIO_PATH):
    root = ET.parse(path).getroot()
    cells = [e for e in root.iter() if e.tag.split("}")[-1] == "mxCell"]
    # mxCell id=0 + id=1 + title + header row (label + months) + data grid
    expected = 2 + 1 + (1 + len(MONTHS)) + len(ACTIVITIES) * (1 + len(MONTHS))
    bars = [e for e in cells if BAR_FILL in (e.get("style") or "")]
    ok = len(cells) == expected and len(bars) == active_cell_count()
    print("drawio : cells=%d (expected %d), bars=%d (expected %d) -> %s"
          % (len(cells), expected, len(bars), active_cell_count(),
             "PASS" if ok else "FAIL"))
    return ok


def verify_docx(path=DOCX_PATH):
    from docx import Document
    doc = Document(path)
    problems = []
    if len(doc.tables) != 2:
        problems.append("expected 2 tables (band + chart), found %d"
                        % len(doc.tables))
    else:
        t = doc.tables[1]
        if len(t.rows) != 1 + len(ACTIVITIES):
            problems.append("rows=%d (expected %d)"
                            % (len(t.rows), 1 + len(ACTIVITIES)))
        if len(t.columns) != 1 + len(MONTHS):
            problems.append("cols=%d (expected %d)"
                            % (len(t.columns), 1 + len(MONTHS)))
        if [t.cell(0, j + 1).text for j in range(len(MONTHS))] != MONTHS:
            problems.append("month header row mismatch")
        shaded = 0
        for i, (label, months) in enumerate(ACTIVITIES, start=1):
            for j, m in enumerate(MONTHS):
                fill = _cell_fill(t.cell(i, j + 1))
                is_bar = bool(fill) and fill.upper() == BAR_FILL
                if is_bar:
                    shaded += 1
                if is_bar != (m in months):
                    problems.append("cell mismatch row %d, %s" % (i, m))
        if shaded != active_cell_count():
            problems.append("shaded=%d (expected %d)"
                            % (shaded, active_cell_count()))
    ok = not problems
    print("docx   : %s -> %s"
          % ("; ".join(problems) if problems else "grid + shading + header OK",
             "PASS" if ok else "FAIL"))
    return ok


def verify_png(path=PNG_PATH):
    from PIL import Image
    img = Image.open(path).convert("RGB")
    bar = tuple(int(BAR_FILL[k:k + 2], 16) for k in (0, 2, 4))
    bad = []
    for i, (label, months) in enumerate(ACTIVITIES):
        for m in MONTHS:
            x, y = png_cell_center(i, m)
            px = img.getpixel((int(x), int(y)))
            if (px == bar) != (m in months):
                bad.append((label[:30], m, px))
    ok = not bad
    print("png    : %d cell probes, %d mismatches -> %s"
          % (len(ACTIVITIES) * len(MONTHS), len(bad), "PASS" if ok else "FAIL"))
    for b in bad[:10]:
        print("   mismatch: %s %s px=%s" % b)
    return ok


def verify_md(path=MD_PATH):
    ok = os.path.exists(path)
    if ok:
        with open(path, encoding="utf-8") as f:
            content = f.read()
        ok = content.startswith("#") and FIG_CAPTION in content
    print("md     : caption + narrative present -> %s" % ("PASS" if ok else "FAIL"))
    return ok


def main():
    generated = [build_drawio(), build_docx(), build_png(), build_md()]
    print("Generated:")
    for p in generated:
        print("  " + p)
    if "verify" in sys.argv:
        print("-" * 70)
        results = [verify_drawio(), verify_docx(), verify_png(), verify_md()]
        print("-" * 70)
        print("VERIFY: %s" % ("ALL PASS" if all(results) else "FAILURES PRESENT"))
        sys.exit(0 if all(results) else 1)


if __name__ == "__main__":
    main()
